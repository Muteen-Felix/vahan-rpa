"""Vahan RPA PoC — Role 3. Selector gắn từ bàn giao Role 2 (mục 4.3 báo cáo).
Attended RPA: người chỉ tương tác TRONG BROWSER (đọc + gõ CAPTCHA), KHÔNG cần
quay lại terminal gõ Enter — Playwright tự poll DOM (wait_for_function) để biết
khi nào người đã gõ xong, rồi tự tiếp quản Apply -> Export -> verify."""

import sys
import time
import zipfile
import json

# [FACT] Fix: chạy qua pipe/redirect trên Windows, stdout mặc định là cp1252 và
# crash khi print tiếng Việt có dấu (UnicodeEncodeError). Ép lại UTF-8 cho an toàn.
if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")

from openpyxl import load_workbook
from playwright.sync_api import sync_playwright, TimeoutError as PWTimeout

from config import DIAGNOSTIC_DIR, DOWNLOAD_DIR
from ui_contract import (
    UIDriftError,
    assert_ui_contract,
    find_all_checkbox,
    find_dropdown_option,
    format_ui_drift,
    get_dropdown_container,
)

URL = "https://analytics.parivahan.gov.in/analytics/vahanpublicreport?lang=en"

# [FACT] Nguồn: bao-cao-vahan-rpa-poc_role2.md mục 4.3 (selector map Role 2 bàn giao),
# đo trực tiếp trên DOM, ID tĩnh (không phải j_idt...). #externalCaptcha cũng có
# trong danh sách ID tĩnh của role2 (mục 4.2).
#
# [FACT] Đã chạy smoke test (headless, screenshot before/after) với select_option
# (force=True) thẳng trên #stateName/#vehicleCategoryGroup/#vehicleFuel: lệnh KHÔNG
# lỗi, nhưng UI vẫn hiện nguyên placeholder "--- Select State/Category/Fuel ---" —
# tức widget hiển thị (custom multiselect-dropdown, JS riêng) KHÔNG đồng bộ theo giá
# trị vừa set trên <select> ẩn. Vì vậy Category/Fuel dùng container-click thật
# (verify bằng assert checkbox.is_checked()), không dùng select_option() thẳng.
#
# [FACT] Y-Axis/X-Axis (#yAxis/#xAxis) là <select> thật, HIỂN THỊ bình thường —
# select_option() chạy thẳng được. Đọc thẳng script nhúng trong trang (không đoán):
# X-Axis CHỈ được JS populate khi có sự kiện "click" trên #yAxis — KHÔNG phải
# "change". Bắt buộc: select_option(yAxis) trước, dispatch click trên #yAxis, rồi
# mới select_option(xAxis) — verify bằng #yAxis_hidden/#xAxis_hidden khớp đúng.
#
# [FACT] Đã test CAPTCHA sai (điền "XXXXXX", headless): trang hiện "Invalid CAPTCHA."
# nhưng KHÔNG reload toàn trang — url không đổi, Category/Fuel/Y-Axis/X-Axis vẫn giữ
# nguyên giá trị đã chọn, chỉ #externalCaptcha bị xoá trắng. Nên khi retry KHÔNG cần
# chọn lại filter, chỉ cần chờ người gõ CAPTCHA mới.
SELECTORS = {
    "state_select": "#stateName",
    "category_select": "#vehicleCategoryGroup",
    "fuel_select": "#vehicleFuel",
    "yaxis": "#yAxis",
    "xaxis": "#xAxis",
    "captcha_input": "#externalCaptcha",
    "apply_button": "#applyTrigger",
    "download_excel_button": "#downloadBtn1",
}

# [FACT] Quan sát trực tiếp (nhiều lần load): CAPTCHA trang này luôn 6 ký tự
# (vd "GDX4F3", "T6SB8u", "a27GR6"). Dùng làm ngưỡng "người đã gõ xong".
CAPTCHA_LENGTH = 6
CAPTCHA_WAIT_TIMEOUT_MS = 300_000  # 5 phút mỗi lượt — đủ cho người đọc + gõ tay


def select_checkbox_option(page, container_selector: str, option_text: str, label: str):
    """Checkbox multi-select searchable dropdown (State/Category/Fuel) — verify thật
    bằng test_category_and_fuel.py. Dùng cho một giá trị cụ thể, KHÔNG dùng cho case "All".

    ``container_selector`` is retained in the function signature for compatibility,
    but callers now pass the underlying select id so the wrapper is resolved by the
    UI contract adapter instead of a document-wide positional XPath.
    """
    hidden_select_id = container_selector.removeprefix("#")
    container = get_dropdown_container(page, hidden_select_id, step=label)
    container.click()

    search_candidates = container.locator(
        "input.multiselect-dropdown-search, input[type='search'], input[placeholder*='search' i]"
    )
    if search_candidates.count() != 1:
        raise UIDriftError(
            "UI_DRIFT_SEARCH_INPUT",
            f"Không tìm thấy duy nhất ô tìm kiếm cho {label}, count={search_candidates.count()}.",
            step=label,
            details={"label": label, "count": search_candidates.count()},
        )
    search_box = search_candidates.first
    search_box.fill(option_text)
    container.locator("[data-search-text]").first.wait_for(state="visible", timeout=5_000)

    option = find_dropdown_option(container, option_text, label=label, step=label)
    option.scroll_into_view_if_needed()
    option.click()

    checkbox = option.locator("input[type='checkbox']")
    if checkbox.count() != 1:
        raise UIDriftError(
            "UI_DRIFT_OPTION_CONTROL",
            f"Option '{option_text}' trong {label} không có đúng một checkbox.",
            step=label,
            details={"label": label, "option": option_text, "checkbox_count": checkbox.count()},
        )
    checkbox = checkbox.first
    assert checkbox.is_checked(), f"[{label}] click xong nhưng checkbox KHÔNG được tick"

    selected = page.locator(f"#{hidden_select_id} option:checked")
    selected_text = [value.strip().casefold() for value in selected.all_text_contents()]
    if option_text.strip().casefold() not in selected_text:
        raise UIDriftError(
            "UI_DRIFT_SELECTION_NOT_SYNCED",
            f"{label} hiển thị đã chọn nhưng select gốc không đồng bộ với '{option_text}'.",
            step=label,
            details={"label": label, "option": option_text, "selected": selected_text},
        )

    page.click("body", position={"x": 2, "y": 2})


def select_all_checkbox(page, container_selector: str, label: str):
    """Riêng case chọn "All" — không phải option thường (không có data-search-text),
    mà là div.multiselect-dropdown-all-selector nằm đầu danh sách. Verify bằng
    inspect_fuel_options.py. KHÔNG dùng select_checkbox_option() cho case này."""
    hidden_select_id = container_selector.removeprefix("#")
    container = get_dropdown_container(page, hidden_select_id, step=label)
    container.click()

    all_checkbox = find_all_checkbox(container, label=label, step=label)
    all_checkbox.click()
    assert all_checkbox.is_checked(), f"[{label}] click xong nhưng checkbox 'All' KHÔNG được tick"

    option_count = page.locator(f"#{hidden_select_id} option").count()
    selected_count = page.locator(f"#{hidden_select_id} option:checked").count()
    if option_count == 0 or selected_count != option_count:
        raise UIDriftError(
            "UI_DRIFT_SELECT_ALL_NOT_SYNCED",
            f"{label} hiển thị All nhưng select gốc không chọn đủ option.",
            step=label,
            details={"label": label, "option_count": option_count, "selected_count": selected_count},
        )

    page.click("body", position={"x": 2, "y": 2})


def apply_filters(page):
    """Category Group = Two Wheeler, Fuel = All, Y-Axis = Fuel, X-Axis = Vehicle
    Category Group. [ASSUMPTION] Không chọn State/Year (dùng mặc định trang) — theo
    yêu cầu phạm vi hiện tại, chưa phải quyết định chính thức mục 0.1 báo cáo."""
    select_checkbox_option(page, SELECTORS["category_select"], "TWO WHEELER", "Category Group")
    # [FACT] mục 0.1 báo cáo (role1) đã chốt Fuel = "All".
    select_all_checkbox(page, SELECTORS["fuel_select"], "Fuel")

    page.select_option(SELECTORS["yaxis"], label="Fuel")
    # Bắt buộc dispatch click để trigger updateXAxisOptions() — xem ghi chú ở SELECTORS.
    page.locator(SELECTORS["yaxis"]).evaluate('el => el.dispatchEvent(new Event("click", {bubbles:true}))')
    page.wait_for_function(
        """
        () => Array.from(document.querySelectorAll('#xAxis option')).some((option) =>
            (option.label || option.textContent || '').trim().toUpperCase() === 'VEHICLE CATEGORY GROUP'
        )
        """,
        timeout=5_000,
    )
    page.select_option(SELECTORS["xaxis"], label="Vehicle Category Group")

    expected_hidden = {
        "#yAxis_hidden": "vehicleFuel",
        "#xAxis_hidden": "vehicleCategoryGroup",
    }
    for selector, expected in expected_hidden.items():
        actual = page.locator(selector).input_value()
        if actual != expected:
            raise UIDriftError(
                "UI_DRIFT_AXIS_NOT_SYNCED",
                f"{selector} không đồng bộ: expected={expected!r}, actual={actual!r}.",
                step="axis",
                details={"selector": selector, "expected": expected, "actual": actual},
            )
    # Tự động focus vào ô CAPTCHA ngay sau khi chọn X-Axis để người dùng gõ được luôn
    captcha_box = page.locator(SELECTORS["captcha_input"])
    captcha_box.scroll_into_view_if_needed()
    captcha_box.focus()


def wait_for_captcha_typed(page, timeout_ms=CAPTCHA_WAIT_TIMEOUT_MS):
    """Attended, KHÔNG dùng input()/terminal. Poll trực tiếp DOM #externalCaptcha —
    coi là "người gõ xong" khi đủ CAPTCHA_LENGTH ký tự. Script tự tiếp quản ngay khi
    điều kiện đúng, không cần người bấm gì thêm ở terminal."""
    captcha_box = page.locator(SELECTORS["captcha_input"])
    captcha_box.scroll_into_view_if_needed()
    captcha_box.focus()
    print(
        f"    >>> Đã focus vào ô CAPTCHA. Đang chờ người đọc CAPTCHA trên browser và gõ đủ {CAPTCHA_LENGTH} ký tự "
        f"vào ô CAPTCHA (tối đa {timeout_ms / 1000:.0f}s)..."
    )
    page.wait_for_function(
        f"""
        () => {{
            const el = document.querySelector('{SELECTORS["captcha_input"]}');
            return !!el && el.value.trim().length >= {CAPTCHA_LENGTH};
        }}
        """,
        timeout=timeout_ms,
    )
    print("    >>> Đã phát hiện CAPTCHA được điền — robot tiếp quản, bấm Apply.")


def verify_file(path):
    """Trả về dict mô tả file tải về, để đối chiếu với baseline ở mục 6."""
    result = {"path": str(path), "exists": path.exists()}
    if not result["exists"]:
        return result

    result["size_bytes"] = path.stat().st_size
    result["is_real_xlsx"] = zipfile.is_zipfile(path)
    if not result["is_real_xlsx"]:
        result["head"] = path.read_bytes()[:200]
        return result

    ws = load_workbook(path, read_only=True).active
    rows = [[c.value for c in r] for r in ws.iter_rows()]
    result["sheet"] = ws.title
    result["row_count_raw"] = len(rows)
    result["columns"] = rows[2] if len(rows) > 2 else None
    result["rows"] = rows
    return result


def validate_report_file(verification: dict) -> None:
    """Fail closed if the downloaded workbook is not the expected report shape."""

    if not verification.get("exists"):
        raise ValueError("File Excel không tồn tại sau khi tải.")
    if not verification.get("is_real_xlsx"):
        raise ValueError("File tải về không phải XLSX hợp lệ.")

    expected_columns = ["Fuel", "Two Wheeler", "Total"]
    actual_columns = verification.get("columns")
    if actual_columns != expected_columns:
        raise ValueError(
            f"Schema file không khớp: expected={expected_columns!r}, actual={actual_columns!r}."
        )


def write_ui_diagnostic(result: dict, error: UIDriftError) -> str:
    """Write metadata-only diagnostics; never persist CAPTCHA or page HTML."""

    path = DIAGNOSTIC_DIR / f"ui-drift-{int(time.time())}.json"
    notification = format_ui_drift(error)
    payload = {
        "iteration": result.get("iteration"),
        "timestamp_epoch": int(time.time()),
        "error_code": error.code,
        "step": error.step,
        "message": str(error),
        "details": error.details,
        "user_notification": notification,
        "ui_contract": result.get("ui_contract"),
    }
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    return str(path)


def run_once(iteration_label="", max_captcha_attempts=3):
    """1 lần chạy đầy đủ: mở trang -> filter -> chờ CAPTCHA (attended) -> Apply ->
    chờ bảng render -> Export -> verify. Trả về dict kết quả + thời gian đo được,
    để ghi trực tiếp vào mục 5.2/5.3/7 báo cáo."""
    t0 = time.time()
    result = {"iteration": iteration_label}
    with sync_playwright() as p:
        # headless=False + maximized: người chạy cần tự mắt thấy trang và gõ CAPTCHA tay
        # (cùng pattern đã dùng ở test_category_and_fuel.py / test_category_dropdown.py).
        browser = p.chromium.launch(headless=False, args=["--start-maximized"])
        context = browser.new_context(accept_downloads=True, no_viewport=True)
        page = context.new_page()
        try:
            print(f"\n=== {iteration_label} ===")
            print("[1/5] Mở trang...")
            t_step = time.time()
            page.goto(URL, wait_until="domcontentloaded")
            result["t_load_page_s"] = round(time.time() - t_step, 1)

            print("[1.5/5] Kiểm tra UI contract trước khi thao tác...")
            result["ui_contract"] = assert_ui_contract(page, step="preflight")

            print("[2/5] Chọn filter (Category=Two Wheeler, Fuel=All, Y/X-Axis)...")
            t_step = time.time()
            apply_filters(page)
            result["t_apply_filters_s"] = round(time.time() - t_step, 1)

            print("[3/5] Chờ CAPTCHA (attended) + Apply, retry nếu sai...")
            t_step = time.time()
            applied = False
            captcha_attempts_used = 0
            for attempt in range(1, max_captcha_attempts + 1):
                captcha_attempts_used = attempt
                print(f"    -- lần thử CAPTCHA {attempt}/{max_captcha_attempts} --")
                try:
                    wait_for_captcha_typed(page)
                except PWTimeout:
                    result["error"] = f"Hết {CAPTCHA_WAIT_TIMEOUT_MS/1000:.0f}s chờ người gõ CAPTCHA, không thấy nhập."
                    break

                page.click(SELECTORS["apply_button"])
                page.wait_for_load_state("networkidle")

                # Apply có thể render lại toàn bộ trang.  Xác nhận contract vẫn
                # giữ nguyên trước khi đọc lỗi CAPTCHA hoặc chạm vào kết quả.
                assert_ui_contract(
                    page,
                    step="post-apply",
                    expected_signature=result["ui_contract"]["signature"],
                )

                if page.get_by_text("Invalid CAPTCHA", exact=False).count() > 0:
                    print(f"    [CAPTCHA lần {attempt}] Sai — trang tự sinh CAPTCHA mới (filter vẫn giữ nguyên), thử lại.")
                    continue

                applied = True
                break

            result["t_captcha_apply_s"] = round(time.time() - t_step, 1)
            result["captcha_attempts_used"] = captcha_attempts_used

            if not applied:
                result.setdefault("error", f"Không vượt qua CAPTCHA sau {max_captcha_attempts} lần")
                result["success"] = False
                return result

            print("[4/5] Chờ bảng dữ liệu render xong (proxy: nút Export xuất hiện)...")
            t_step = time.time()
            page.wait_for_selector(SELECTORS["download_excel_button"], state="visible", timeout=30_000)
            result["t_wait_table_s"] = round(time.time() - t_step, 1)

            print("[5/5] Bấm Export, tải file...")
            t_step = time.time()
            with page.expect_download() as dl:
                page.click(SELECTORS["download_excel_button"])
            # Thêm timestamp vào tên file — tránh 3 lần chạy lặp ghi đè lẫn nhau,
            # cần giữ đủ cả 3 file để Role 1 đối chiếu mục 6.
            stamped_name = f"{int(t0)}_{dl.value.suggested_filename}"
            path = DOWNLOAD_DIR / stamped_name
            dl.value.save_as(path)
            result["t_download_s"] = round(time.time() - t_step, 1)

            verify = verify_file(path)
            result.update(verify)
            validate_report_file(verify)
            result["success"] = True
            return result

        except UIDriftError as e:
            result["success"] = False
            result["ui_drift"] = True
            result["error_code"] = e.code
            notification = format_ui_drift(e)
            result["error"] = notification["message"]
            result["ui_drift_notification"] = notification
            result["diagnostic_path"] = write_ui_diagnostic(result, e)
            return result
        except PWTimeout as e:
            result["success"] = False
            result["error"] = f"TIMEOUT: {e}"
            return result
        except Exception as e:
            result["success"] = False
            result["error"] = f"{type(e).__name__}: {e}"
            return result
        finally:
            result["duration_s"] = round(time.time() - t0, 1)
            browser.close()


def run_iterations(n=3):
    """Khối 4: chạy lặp n lần KHÔNG sửa code giữa các lần, đo thời gian + lỗi từng
    lần cho mục 5.3. Mỗi lần vẫn cần người đọc CAPTCHA trên browser (không có cách
    unattended thật, xem mục 9.2 giới hạn của PoC)."""
    results = []
    for i in range(1, n + 1):
        r = run_once(iteration_label=f"Lần {i}/{n}")
        results.append(r)
        print(f"--- Kết quả lần {i}: success={r.get('success')} duration_s={r.get('duration_s')} "
              f"error={r.get('error')} ---")

    success_count = sum(1 for r in results if r.get("success"))
    print(f"\n=== TỔNG KẾT: {success_count}/{n} lần thành công ===")
    for r in results:
        print(r)
    return results


if __name__ == "__main__":
    run_iterations(n=3)
